# app.py
# -----------------------------------------------------------------------------
# AI Resume Optimizer (no external APIs)
# - Streamlit UI
# - PDF/DOCX/TXT extraction
# - Heuristic analyzer returning ATS-style results
# -----------------------------------------------------------------------------

from __future__ import annotations

import io
import re
from enum import Enum
from typing import Any, Dict, List, Tuple

import pandas as pd
import streamlit as st
import plotly.graph_objects as go

# Optional docx support (already in your env per log; fallback safe)
try:
    import docx  # python-docx
except Exception:
    docx = None


# =========================
# Domain: Enums & Utilities
# =========================
class JobField(str, Enum):
    DATA_SCIENCE = "Data Science"
    MACHINE_LEARNING = "Machine Learning"
    SOFTWARE_ENGINEERING = "Software Engineering"
    DEVOPS = "DevOps"
    GENERATIVE_AI = "Generative AI"
    CYBERSECURITY = "Cybersecurity"
    PRODUCT_MANAGEMENT = "Product Management"
    OTHER = "Other"


class FileHandler:
    MAX_MB = 10
    ALLOWED = {"pdf", "docx", "txt"}

    @staticmethod
    def validate_file(uploaded_file) -> Tuple[bool, str]:
        name = getattr(uploaded_file, "name", "file")
        suffix = name.split(".")[-1].lower()
        if suffix not in FileHandler.ALLOWED:
            return False, f"Unsupported type: .{suffix}. Allowed: {', '.join(FileHandler.ALLOWED)}"
        size = getattr(uploaded_file, "size", None)
        if size is not None and size > FileHandler.MAX_MB * 1024 * 1024:
            return False, f"File too large. Max {FileHandler.MAX_MB} MB."
        return True, f"✅ {name} uploaded (.{suffix})."

    @staticmethod
    def extract_text(uploaded_file) -> str:
        name = getattr(uploaded_file, "name", "file")
        suffix = name.split(".")[-1].lower()
        data = uploaded_file.read()

        # Reset pointer for potential re-reads by Streamlit
        try:
            uploaded_file.seek(0)
        except Exception:
            pass

        if suffix == "txt":
            try:
                return data.decode("utf-8", errors="ignore")
            except Exception:
                return data.decode("latin-1", errors="ignore")

        if suffix == "docx":
            if docx is None:
                return ""
            f = io.BytesIO(data)
            d = docx.Document(f)
            return "\n".join(p.text for p in d.paragraphs)

        if suffix == "pdf":
            # Try pdfminer; fallback to pdfplumber; else empty
            try:
                import pdfminer.high_level as pm
                return pm.extract_text(io.BytesIO(data)) or ""
            except Exception:
                try:
                    import pdfplumber
                    text = []
                    with pdfplumber.open(io.BytesIO(data)) as pdf:
                        for page in pdf.pages:
                            text.append(page.extract_text() or "")
                    return "\n".join(text)
                except Exception:
                    return ""
        return ""


# =========================
# Heuristic Analyzer (local)
# =========================
class EnhancedResumeOptimizationCrew:
    """
    Heuristic resume analyzer that returns the structure the UI expects.
    Replace with your real CrewAI/LLM pipeline later if you want.
    """

    def __init__(self, job_field: str) -> None:
        self.job_field = job_field

    def analyze_resume(self, text: str) -> Dict[str, Any]:
        text_lower = text.lower()

        # Detect common sections
        section_patterns = {
            "summary": r"(?:^|\n)\s*(summary|objective)\s*[:\n]",
            "experience": r"(?:^|\n)\s*(experience|work history)\s*[:\n]",
            "education": r"(?:^|\n)\s*(education|coursework)\s*[:\n]",
            "projects": r"(?:^|\n)\s*(projects|project experience)\s*[:\n]",
            "skills": r"(?:^|\n)\s*(skills|technical skills)\s*[:\n]",
            "certifications": r"(?:^|\n)\s*(certifications|licenses)\s*[:\n]",
        }
        sections_found: Dict[str, str] = {}
        for name, pat in section_patterns.items():
            m = re.search(pat, text_lower)
            if m:
                start = m.start()
                # naive next header heuristic
                next_hdr = re.search(r"\n[A-Z][A-Za-z0-9 &+/]{2,}\n", text[m.end():])
                end = m.end() + (next_hdr.start() if next_hdr else 600)
                snippet = text[start:end].strip()
                sections_found[name] = snippet

        # Field-specific keywords (toy lists)
        field_keywords = {
            "Data Science": ["python", "pandas", "numpy", "sklearn", "sql"],
            "Machine Learning": ["ml", "pytorch", "tensorflow", "xgboost", "model"],
            "Software Engineering": ["microservices", "rest", "docker", "kubernetes", "typescript", "java"],
            "DevOps": ["ci/cd", "kubernetes", "docker", "terraform", "aws"],
            "Generative AI": ["langchain", "rag", "prompt", "openai", "llm", "vector", "embedding"],
            "Cybersecurity": ["owasp", "siem", "ids", "iam", "zero trust"],
            "Product Management": ["roadmap", "kpi", "user research", "backlog", "stakeholder"],
            "Other": [],
        }
        kw_list = field_keywords.get(self.job_field, [])
        words = re.findall(r"\b[\w\-]+\b", text_lower)
        total_words = max(len(words), 1)

        keyword_analysis: Dict[str, Dict[str, float | int]] = {}
        hit_count = 0
        for k in kw_list:
            c = text_lower.count(k.lower())
            if c > 0:
                hit_count += 1
            keyword_analysis[k] = {"count": c, "density": c / total_words}

        # Readability approximations
        sentences = re.split(r"[.!?]+", text)
        sent_list = [s for s in sentences if s.strip()]
        sentence_count = max(len(sent_list), 1)
        avg_sentence_len = len(words) / sentence_count

        # Scoring heuristics
        section_score = int(min(100, 20 + len(sections_found) * 12))
        keyword_score = int(min(100, 40 + hit_count * 8))
        # crude "format" cue
        format_score = 85 if ("•" in text or "—" in text or " - " in text) else 70
        readability_score = int(max(50, min(100, 100 - abs(avg_sentence_len - 18) * 5)))
        completeness = int(min(100, (section_score + keyword_score + format_score + readability_score) / 4))
        ats_total = int(round(0.35 * keyword_score + 0.25 * section_score + 0.20 * format_score + 0.20 * readability_score))

        # Contact info
        contact: Dict[str, str] = {}
        email = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)
        phone = re.search(r"(\+?\d[\d\s\-().]{7,}\d)", text)
        if email:
            contact["email"] = email.group(0)
        if phone:
            contact["phone"] = phone.group(0)

        # Recommendations
        recs: List[str] = []
        if hit_count < max(3, len(kw_list) // 2):
            recs.append(f"Add more {self.job_field} keywords (found {hit_count}/{len(kw_list)} core terms).")
        if "summary" not in sections_found:
            recs.append("Add a concise 2–3 line Summary at the top with role + impact metrics.")
        if avg_sentence_len > 24:
            recs.append("Shorten long sentences; aim for ~15–20 words per sentence.")
        if "skills" not in sections_found:
            recs.append("Include a Skills section with grouped tools (e.g., Languages, Frameworks, Cloud).")
        if not recs:
            recs.append("Looks strong. Consider adding quantified impact bullets (%, time saved, cost reduced).")

        return {
            "success": True,
            "ats_scores": {
                "total_score": ats_total,
                "keyword_relevance": keyword_score,
                "section_structure": section_score,
                "format_compatibility": format_score,
                "readability": readability_score,
                "completeness": completeness,
            },
            "sections": sections_found,
            "keyword_analysis": keyword_analysis,
            "readability": {
                "word_count": len(words),
                "sentence_count": sentence_count,
                "avg_sentence_length": avg_sentence_len,
                "flesch_reading_ease": 60.0,        # placeholders
                "flesch_kincaid_grade": 10.0,       # placeholders
                "automated_readability_index": 9.5, # placeholders
            },
            "recommendations": recs,
            "contact_info": contact,
        }


# =========================
# UI: Results Renderer
# =========================
def display_results(results: Dict[str, Any]) -> None:
    """Display analysis results in the right column"""
    st.header("📊 Analysis Results")

    # ATS Compatibility Score
    scores = results.get("ats_scores", {})
    total_score = float(scores.get("total_score", 0))

    # Gauge
    fig = go.Figure(
        go.Indicator(
            mode="gauge+number+delta",
            value=total_score,
            domain={"x": [0, 1], "y": [0, 1]},
            title={"text": "ATS Compatibility Score"},
            delta={"reference": 80},
            gauge={
                "axis": {"range": [None, 100]},
                "bar": {"color": "darkblue"},
                "steps": [
                    {"range": [0, 50], "color": "lightgray"},
                    {"range": [50, 80], "color": "yellow"},
                    {"range": [80, 100], "color": "green"},
                ],
                "threshold": {"line": {"color": "red", "width": 4}, "thickness": 0.75, "value": 90},
            },
        )
    )
    st.plotly_chart(fig, use_container_width=True)

    # Score breakdown
    c1, c2, c3 = st.columns(3)
    with c1:
        st.metric("Keyword Relevance", f"{float(scores.get('keyword_relevance', 0)):.1f}%")
        st.metric("Section Structure", f"{float(scores.get('section_structure', 0)):.1f}%")
    with c2:
        st.metric("Format Compatibility", f"{float(scores.get('format_compatibility', 0)):.1f}%")
        st.metric("Readability", f"{float(scores.get('readability', 0)):.1f}%")
    with c3:
        st.metric("Completeness", f"{float(scores.get('completeness', 0)):.1f}%")

    # Tabs
    tab1, tab2, tab3, tab4 = st.tabs(["📝 Sections", "🔑 Keywords", "📊 Readability", "💡 Recommendations"])

    with tab1:
        st.subheader("Resume Sections Analysis")
        sections = results.get("sections", {})
        if sections:
            for section_name, content in sections.items():
                with st.expander(f"{section_name.title()} Section"):
                    st.text_area("Content", content, height=120, disabled=True)
        else:
            st.warning("No standard sections detected. Consider organizing your resume with clear section headers.")

    with tab2:
        st.subheader("Keyword Analysis")
        keyword_analysis = results.get("keyword_analysis", {})
        if keyword_analysis:
            keyword_data = []
            for keyword, data in keyword_analysis.items():
                keyword_data.append(
                    {
                        "Keyword": keyword,
                        "Count": int(data["count"]),
                        "Density": f"{float(data['density']) * 100:.2f}%",
                    }
                )
            df = pd.DataFrame(keyword_data)
            st.dataframe(df, use_container_width=True)
        else:
            st.info("No keyword analysis available for the selected field.")

    with tab3:
        st.subheader("Readability Metrics")
        readability = results.get("readability", {})
        if readability:
            cc1, cc2 = st.columns(2)
            with cc1:
                st.metric("Word Count", int(readability.get("word_count", 0)))
                st.metric("Sentence Count", int(readability.get("sentence_count", 0)))
                st.metric("Avg Sentence Length", f"{float(readability.get('avg_sentence_length', 0)):.1f}")
            with cc2:
                st.metric("Flesch Reading Ease", f"{float(readability.get('flesch_reading_ease', 0)):.1f}")
                st.metric("Flesch-Kincaid Grade", f"{float(readability.get('flesch_kincaid_grade', 0)):.1f}")
                st.metric(
                    "Automated Readability Index",
                    f"{float(readability.get('automated_readability_index', 0)):.1f}",
                )
        else:
            st.info("No readability metrics available.")

    with tab4:
        st.subheader("Optimization Recommendations")
        recommendations = results.get("recommendations", [])
        if recommendations:
            for i, rec in enumerate(recommendations, 1):
                st.write(f"{i}. {rec}")
        else:
            st.success("Great job! No major improvements needed.")

        contact_info = results.get("contact_info", {})
        if contact_info:
            st.subheader("Contact Information Found")
            for key, value in contact_info.items():
                st.write(f"**{key.title()}:** {value}")


# =========================
# UI: Main
# =========================
def main() -> None:
    st.set_page_config(page_title="AI Resume Optimizer", page_icon="📄", layout="wide")

    st.title("🚀 AI-Powered Resume Optimizer")
    st.markdown(
        "Upload your resume and get comprehensive ATS compatibility analysis with actionable insights."
    )

    # Sidebar
    with st.sidebar:
        st.header("Configuration")

        job_field = st.selectbox(
            "Select Job Field",
            options=[field.value for field in JobField],
            index=0,
        )

        analysis_type = st.selectbox(
            "Analysis Type",
            options=["Comprehensive", "Quick", "ATS Focus"],
            index=0,
        )

        st.header("About")
        st.info(
            "This tool analyzes your resume for ATS compatibility, keyword optimization, and professional formatting."
        )

    # Main layout
    left, right = st.columns([1, 2])

    with left:
        st.header("Upload Resume")
        uploaded_file = st.file_uploader(
            "Choose your resume file",
            type=["pdf", "docx", "txt"],
            help="Upload a PDF, DOCX, or TXT file (max 10MB)",
        )

        if uploaded_file:
            is_valid, validation_message = FileHandler.validate_file(uploaded_file)
            if is_valid:
                st.success(validation_message)

                if st.button("Analyze Resume", type="primary"):
                    with st.spinner("Analyzing your resume..."):
                        resume_text = FileHandler.extract_text(uploaded_file)
                        if resume_text:
                            analyzer = EnhancedResumeOptimizationCrew(job_field=job_field)
                            results = analyzer.analyze_resume(resume_text)

                            if results.get("success"):
                                st.session_state["analysis_results"] = results
                                st.session_state["resume_text"] = resume_text
                                st.rerun()
                            else:
                                st.error(f"Analysis failed: {results.get('error', 'Unknown error')}")
                        else:
                            st.error("Could not extract text from the uploaded file.")
            else:
                st.error(validation_message)

    with right:
        if "analysis_results" in st.session_state:
            display_results(st.session_state["analysis_results"])
        else:
            st.header("Analysis Results")
            st.info("Upload a resume and click 'Analyze Resume' to see results here.")


if __name__ == "__main__":
    main()
